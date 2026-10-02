"""Document loaders for FuelGuard RAG ingestion pipeline.

Supports Markdown (.md), Plain Text (.txt), JSON (.json), PDF (.pdf),
and DOCX (.docx) with metadata extraction and clean text normalization.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


@dataclass
class LoadedDocument:
    document_id: str
    filename: str
    category: str
    document_type: str
    version: str
    year: int | None
    section: str | None
    page: int | None
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)
    checksum: str = ""


def compute_sha256(file_path: Path) -> str:
    h = hashlib.sha256()
    with file_path.open("rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def _parse_yaml_frontmatter(content: str) -> tuple[dict[str, Any], str]:
    """Extracts frontmatter if present between --- delimiters."""
    meta: dict[str, Any] = {}
    clean_text = content
    if content.startswith("---"):
        parts = content.split("---", 2)
        if len(parts) >= 3:
            raw_meta = parts[1]
            clean_text = parts[2].strip()
            # Simple line-by-line YAML parser to avoid hard dependency on PyYAML if minimal
            for line in raw_meta.splitlines():
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                if ":" in line:
                    key, val = line.split(":", 1)
                    key = key.strip()
                    val = val.strip().strip('"').strip("'")
                    if val.isdigit():
                        meta[key] = int(val)
                    else:
                        meta[key] = val
    return meta, clean_text


class DocumentLoader:
    """Multi-format loader for the RAG ingestion pipeline."""

    SUPPORTED_EXTENSIONS = {".md", ".txt", ".json", ".pdf", ".docx"}

    @classmethod
    def load(cls, file_path: str | Path, default_category: str | None = None) -> list[LoadedDocument]:
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")

        ext = path.suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            return []

        checksum = compute_sha256(path)
        category = default_category or path.parent.name
        filename = path.name

        if ext in (".md", ".txt"):
            return cls._load_text_or_markdown(path, checksum, category, filename)
        elif ext == ".json":
            return cls._load_json(path, checksum, category, filename)
        elif ext == ".pdf":
            return cls._load_pdf(path, checksum, category, filename)
        elif ext == ".docx":
            return cls._load_docx(path, checksum, category, filename)
        return []

    @classmethod
    def _load_text_or_markdown(
        cls, path: Path, checksum: str, category: str, filename: str
    ) -> list[LoadedDocument]:
        content = path.read_text(encoding="utf-8", errors="replace")
        frontmatter, clean_text = _parse_yaml_frontmatter(content)

        doc_id = frontmatter.get("document_id") or f"doc-{path.stem}"
        doc_category = frontmatter.get("category") or category
        doc_type = frontmatter.get("document_type") or ("markdown" if path.suffix == ".md" else "text")
        version = str(frontmatter.get("version", "1.0.0"))
        year = frontmatter.get("year")
        section = frontmatter.get("section")

        return [
            LoadedDocument(
                document_id=doc_id,
                filename=filename,
                category=doc_category,
                document_type=doc_type,
                version=version,
                year=int(year) if year and str(year).isdigit() else None,
                section=section,
                page=1,
                text=clean_text,
                metadata={**frontmatter, "file_path": str(path)},
                checksum=checksum,
            )
        ]

    @classmethod
    def _load_json(cls, path: Path, checksum: str, category: str, filename: str) -> list[LoadedDocument]:
        content = path.read_text(encoding="utf-8", errors="replace")
        try:
            data = json.loads(content)
        except Exception:
            return []

        doc_id = f"doc-{path.stem}"
        text_content = json.dumps(data, indent=2)
        return [
            LoadedDocument(
                document_id=doc_id,
                filename=filename,
                category=category,
                document_type="structured_json",
                version="1.0.0",
                year=None,
                section="root",
                page=1,
                text=text_content,
                metadata={"file_path": str(path), "keys": list(data.keys()) if isinstance(data, dict) else []},
                checksum=checksum,
            )
        ]

    @classmethod
    def _load_pdf(cls, path: Path, checksum: str, category: str, filename: str) -> list[LoadedDocument]:
        """PDF extraction using pypdf if installed; fallback to raw byte text extraction."""
        docs: list[LoadedDocument] = []
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(
                        LoadedDocument(
                            document_id=f"doc-{path.stem}-p{i+1}",
                            filename=filename,
                            category=category,
                            document_type="pdf",
                            version="1.0.0",
                            year=None,
                            section=f"page_{i+1}",
                            page=i + 1,
                            text=text.strip(),
                            metadata={"file_path": str(path), "total_pages": len(reader.pages)},
                            checksum=checksum,
                        )
                    )
        except ImportError:
            # Simple text stream scanner if pypdf is not available
            raw = path.read_bytes()
            extracted = re.findall(b"[a-zA-Z0-9.,;:!?'\"/\\-_() \t\n\r]{4,}", raw)
            decoded = " ".join(part.decode("utf-8", "ignore") for part in extracted if len(part) > 10)
            if decoded.strip():
                docs.append(
                    LoadedDocument(
                        document_id=f"doc-{path.stem}",
                        filename=filename,
                        category=category,
                        document_type="pdf",
                        version="1.0.0",
                        year=None,
                        section=None,
                        page=1,
                        text=decoded[:10000],
                        metadata={"file_path": str(path)},
                        checksum=checksum,
                    )
                )
        return docs

    @classmethod
    def _load_docx(cls, path: Path, checksum: str, category: str, filename: str) -> list[LoadedDocument]:
        """DOCX extraction using python-docx if installed, or zipxml extraction."""
        text_lines: list[str] = []
        try:
            import docx
            doc = docx.Document(str(path))
            for p in doc.paragraphs:
                if p.text.strip():
                    text_lines.append(p.text.strip())
        except ImportError:
            import zipfile
            import xml.etree.ElementTree as ET
            try:
                with zipfile.ZipFile(str(path)) as z:
                    xml_content = z.read("word/document.xml")
                    tree = ET.fromstring(xml_content)
                    for node in tree.iter():
                        if node.tag.endswith("t") and node.text:
                            text_lines.append(node.text)
            except Exception:
                pass

        full_text = "\n".join(text_lines).strip()
        if not full_text:
            return []

        return [
            LoadedDocument(
                document_id=f"doc-{path.stem}",
                filename=filename,
                category=category,
                document_type="docx",
                version="1.0.0",
                year=None,
                section=None,
                page=1,
                text=full_text,
                metadata={"file_path": str(path)},
                checksum=checksum,
            )
        ]
